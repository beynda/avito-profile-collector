from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch
import sys
import logging

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import run
from src.config import Config, Region, Delays, BrowserConfig, OutputConfig
from src.db import ProfileRecord
from openpyxl import load_workbook

with TemporaryDirectory() as folder:
    cfg = Config(['sample'], [Region('sample', 'Example City')], Delays(), 1,
                 BrowserConfig(), None, OutputConfig(), 'data/state.db', Path(folder))
    context = Mock()
    context.close.side_effect = KeyboardInterrupt
    playwright = Mock()

    def interrupted_collection(context, cfg, db, logger):
        db.upsert_profile(ProfileRecord('sample', 'https://example.com/profile',
                          'Example Company', 'Example City', '', '',
                          'https://example.com/listing', 'Example listing', 'sample'))
        raise KeyboardInterrupt

    with patch.object(run, 'prepare_session', return_value=True), \
         patch.object(run, 'load_config', return_value=cfg), \
         patch.object(run, 'launch_browser', return_value=(playwright, context)), \
         patch.object(run, 'run_scrape', side_effect=interrupted_collection):
        assert run.main(['--collect']) == 130
    playwright.stop.assert_called_once()
    assert load_workbook(Path(folder) / 'data/output.xlsx').active.max_row == 2
    assert (Path(folder) / 'data/output.csv').exists()
    with patch.object(run, 'launch_browser') as launch:
        assert run.main([]) == 0
        launch.assert_not_called()
    for handler in logging.getLogger('avito_parser').handlers:
        handler.close()
print('PASS: interruption, cleanup interruption, export, safe default')
