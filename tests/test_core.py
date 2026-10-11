from pathlib import Path
import shutil,sys,sqlite3,logging,json,tempfile,ast
from dataclasses import fields
root=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root))
from src.db import open_db,ProfileRecord
from src.export import export
from src import scraper
from src.config import Region
from openpyxl import load_workbook
results={}
for file in root.rglob('*.py'):ast.parse(file.read_text(encoding='utf-8-sig'))
results['syntax']='pass'
assert scraper.extract_profile_id('https://www.avito.ru/user/abcdef?x=1')=='abcdef'
assert 'p=2' in scraper.build_search_url(Region('test','Test'),'roof',2)
results['profile_id_and_search_url']='pass'
used=set()
tree=ast.parse((root/'src/scraper.py').read_text(encoding='utf-8'))
for node in ast.walk(tree):
 if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id in ['first_match_text','first_match_href'] and len(node.args)>1 and isinstance(node.args[1],ast.Constant):used.add(node.args[1].value)
results['missing_selector_keys']=sorted(used-set(scraper.SELECTORS))
with tempfile.TemporaryDirectory() as temp:
 temp=Path(temp)
 with open_db(temp/'state.db') as db:
  values={f.name:'' for f in fields(ProfileRecord)}
  values.update(profile_id='sample',profile_url='https://example.com/profile/sample',name='Sample Company',city='Sample City',email='contact@example.com',listing_url='https://example.com/listing/sample')
  db.upsert_profile(ProfileRecord(**values));db.upsert_profile(ProfileRecord(**values))
  assert db.count_profiles()==1
  db.mark_listing_processed('https://example.com/listing/sample','roof','sample')
  assert db.is_listing_processed('https://example.com/listing/sample')
  export(db,temp/'output.xlsx',temp/'output.csv',logging.getLogger('test'))
  sheet=load_workbook(temp/'output.xlsx').active
  assert sheet.max_row==2 and sheet.max_column==11
  results['sqlite_dedup_resume_csv_xlsx']='pass'
  values.update(profile_id='formula',name='=1+1')
  db.upsert_profile(ProfileRecord(**values))
  export(db,temp/'formula.xlsx',temp/'formula.csv',logging.getLogger('test'))
  sheet=load_workbook(temp/'formula.xlsx').active
  results['xlsx_formula_injection']=any(c.data_type=='f' for row in sheet for c in row)
  assert not results['xlsx_formula_injection']
  try:export(db,temp/'existing/output.xlsx',temp/'different/csv/output.csv',logging.getLogger('test'));results['separate_csv_directory']='pass'
  except Exception as e:results['separate_csv_directory']=type(e).__name__
  assert results['separate_csv_directory']=='pass'
(Path(tempfile.gettempdir())/'avito-public-core-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=True,indent=2))
