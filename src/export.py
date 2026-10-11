from __future__ import annotations
import logging
import re
from pathlib import Path
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .db import Database
COLUMNS = [('name', 'Name / company'), ('profile_url', 'Profile URL'), ('city', 'City'), ('address', 'Address'), ('email', 'Email'), ('listing_url', 'Listing URL'), ('listing_title', 'Listing title'), ('query', 'Search query'), ('collected_at', 'Collected at'), ('listing_count', 'Listing count'), ('description', 'Description')]
COLUMN_WIDTHS = {'Name / company': 26, 'Profile URL': 42, 'City': 20, 'Address': 40, 'Email': 22, 'Listing URL': 46, 'Listing title': 40, 'Search query': 20, 'Collected at': 18, 'Listing count': 14, 'Description': 40}
WS_WHITESPACE_RE = re.compile('\\s+')

def clean_text(value) -> str:
    if value is None:
        return ''
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))
    text = text.replace('\xa0', ' ')
    text = WS_WHITESPACE_RE.sub(' ', text).strip()
    return text

def export(db: Database, excel_path: Path, csv_path: Path, logger: logging.Logger):
    rows = db.all_profiles()
    if not rows:
        logger.warning('No saved profiles; export skipped.')
        return
    data = []
    for row in rows:
        record = {}
        for key, label in COLUMNS:
            value = row[key]
            text = clean_text(value)
            if text.startswith(('=', '+', '-', '@')):
                text = "'" + text
            record[label] = value if key == 'listing_count' else text
        data.append(record)
    df = pd.DataFrame(data, columns=[label for _, label in COLUMNS])
    df.sort_values(by=['City', 'Name / company'], inplace=True, kind='stable')
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        sheet_name = 'Profiles'
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        for idx, (_, label) in enumerate(COLUMNS, start=1):
            col_letter = get_column_letter(idx)
            ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(label, 20)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            ws.row_dimensions[row[0].row].height = 60
    df.to_csv(csv_path, index=False, encoding='utf-8-sig')
    logger.info(f'Exported {len(df)} profiles to {excel_path} and {csv_path}')
